"""
Build the Word (.docx) version of the paper from the shared content blocks.

The prose, tables and figures are defined once in build_paper.py (module-level
`story` list); this script renders the same blocks with python-docx using the
course formatting: 12-point Times New Roman, double spaced, 1-inch margins.

Run:  python3 analyze_demand.py && python3 paper/build_docx.py
Output: paper/gasoline_demand_paper.docx
"""

import html
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_paper as bp                      # noqa: E402  (shared content)

from docx import Document                     # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH # noqa: E402
from docx.oxml import OxmlElement             # noqa: E402
from docx.oxml.ns import qn                   # noqa: E402
from docx.opc.constants import (              # noqa: E402
    RELATIONSHIP_TYPE as RT)
from docx.shared import Inches, Pt            # noqa: E402

OUT = os.path.join(HERE, "gasoline_demand_paper.docx")
FONT = "Times New Roman"
BODY_PT = 12

TAG = re.compile(r'(<b>|</b>|<i>|</i>|<sub>|</sub>|<sup>|</sup>'
                 r'|<link href="[^"]*"[^>]*>|</link>)')


def _font(run, size=BODY_PT, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(attr), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def add_hyperlink(par, url, text, size=BODY_PT):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(attr), FONT)
    rpr.append(rf)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0000FF')
    rpr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rpr.append(u)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_runs(par, text, size=BODY_PT, bold=False, italic=False):
    """Write text into a paragraph, honouring the small markup vocabulary
    shared with the PDF builder (<b>, <i>, <sub>, <link href=...>)."""
    b, i, sub, url = bold, italic, False, None
    for part in TAG.split(text):
        if not part:
            continue
        if part == '<b>':
            b = True
        elif part == '</b>':
            b = bold
        elif part == '<i>':
            i = True
        elif part == '</i>':
            i = italic
        elif part == '<sub>':
            sub = True
        elif part == '</sub>':
            sub = False
        elif part.startswith('<link '):
            m = re.search(r'href="([^"]+)"', part)
            url = m.group(1) if m else None
        elif part == '</link>':
            url = None
        else:
            s = html.unescape(part)
            if url:
                add_hyperlink(par, url, s, size=size)
            else:
                r = _font(par.add_run(s), size=size, bold=b, italic=i)
                if sub:
                    r.font.subscript = True
    return par


def style_par(par, *, double=True, first_indent=None, left_indent=None,
              align=None, before=0, after=0, keep_next=False):
    pf = par.paragraph_format
    pf.line_spacing = 2.0 if double else 1.0
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if first_indent is not None:
        pf.first_line_indent = Inches(first_indent)
    if left_indent is not None:
        pf.left_indent = Inches(left_indent)
    if align is not None:
        par.alignment = align
    pf.keep_with_next = keep_next
    return par


def new_par(doc, text=None, **kw):
    par = doc.add_paragraph()
    style_par(par, **kw)
    if text:
        add_runs(par, text)
    return par


def add_page_number(par):
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    _font(run, size=10)
    for tag, attrs, text in (('w:fldChar', {'w:fldCharType': 'begin'}, None),
                             ('w:instrText', {'xml:space': 'preserve'}, 'PAGE'),
                             ('w:fldChar', {'w:fldCharType': 'end'}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)


def add_table(doc, block):
    rows = bp.table_rows(block)
    size = block["fontsize"]
    widths = block["widths"] or [6.5 / len(rows[0])] * len(rows[0])
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.autofit = False
    numeric_from = block["numeric_from"]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Inches(widths[c] / 72.0)
            par = cell.paragraphs[0]
            style_par(par, double=False, after=0)
            if r == 0:
                add_runs(par, f"<b>{html.escape(val)}</b>", size=size)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                add_runs(par, html.escape(val), size=size)
                if c >= numeric_from:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_par(doc.add_paragraph(), double=False, after=6)
    return t


def build(blocks, path):
    doc = Document()

    # ── page setup: 1-inch margins, letter ───────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── base style: 12-pt Times New Roman, double spaced ─────────────
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(attr), FONT)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    add_page_number(doc.sections[0].footer.paragraphs[0])

    props = doc.core_properties
    props.title = ("The Impact of Gasoline Prices on Gasoline Demand "
                   "in the United States")
    props.author = bp.AUTHOR
    props.subject = "ECON 515 applied research paper"

    for b in blocks:
        k = b["k"]
        if k == "title":
            new_par(doc, b["text"], double=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                    after=6)
            for r in doc.paragraphs[-1].runs:
                r.font.size = Pt(14)
                r.bold = True
        elif k == "sub":
            new_par(doc, b["text"], align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
        elif k == "auth":
            new_par(doc, b["text"], align=WD_ALIGN_PARAGRAPH.CENTER)
        elif k == "spacer":
            style_par(doc.add_paragraph(), double=False, after=b["h"])
        elif k == "h1":
            new_par(doc, f"<b>{html.escape(b['text'])}</b>", first_indent=0,
                    before=12, keep_next=True)
        elif k == "h2":
            new_par(doc, f"<b>{html.escape(b['text'])}</b>", first_indent=0,
                    before=12, keep_next=True)
        elif k == "exh0":
            new_par(doc, f"<b>{html.escape(b['text'])}</b>", first_indent=0,
                    align=WD_ALIGN_PARAGRAPH.CENTER, after=6, keep_next=True)
        elif k == "apph":
            new_par(doc, f"<b>{html.escape(b['text'])}</b>", first_indent=0,
                    align=WD_ALIGN_PARAGRAPH.CENTER, after=12, keep_next=True)
            for r in doc.paragraphs[-1].runs:
                r.font.size = Pt(14)
        elif k == "p":
            new_par(doc, b["text"], first_indent=0.5)
        elif k == "ref":
            new_par(doc, b["text"], first_indent=-0.5, left_indent=0.5)
        elif k == "eq":
            new_par(doc, b["text"], first_indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif k == "exh":
            new_par(doc, f"<b>{html.escape(b['text'])}</b>", first_indent=0,
                    align=WD_ALIGN_PARAGRAPH.CENTER, before=8, keep_next=True)
        elif k == "cap":
            par = new_par(doc, b["text"], double=False, first_indent=0,
                          align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=10)
            for r in par.runs:
                r.italic = True
                r.font.size = Pt(9.5)
        elif k == "code":
            par = new_par(doc, double=False, first_indent=0, left_indent=0.5,
                          before=2, after=2)
            for line in b["lines"]:
                if par.text:
                    par.add_run().add_break()
                r = _font(par.add_run(line), size=10)
                r.font.name = 'Courier'
        elif k == "table":
            add_table(doc, b)
        elif k == "figure":
            img_path = os.path.join(bp.FIG, b["fname"])
            w = min(b["max_width"] / 72.0, 6.5)
            par = new_par(doc, double=False, first_indent=0,
                          align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6)
            par.add_run().add_picture(img_path, width=Inches(w))
        elif k == "pagebreak":
            doc.add_page_break()
        else:
            raise ValueError(f"unknown block kind: {k}")

    doc.save(path)
    return doc


# ── rough page-count estimate for the main text ──────────────────────
def estimate_body_pages(blocks):
    lines = 0.0
    chars_per_line = 95          # Times 12 pt across a 6.5-inch text column
    for b in blocks:
        k = b["k"]
        if k == "h1":
            lines += 1 + 0.5     # heading plus its leading space
        elif k in ("p", "ref"):
            text = re.sub(r'<[^>]+>', '', b["text"])
            lines += math.ceil(len(text) / chars_per_line) + 0.5
        elif k in ("title", "sub", "auth", "spacer", "eq"):
            lines += 1
        elif k == "pagebreak":
            break                # References start a new page
    return lines * 24.0 / (9 * 72)    # 24 pt leading, 9-inch text height


if __name__ == "__main__":
    body_blocks = bp.story
    doc = build(body_blocks, OUT)
    est = estimate_body_pages(body_blocks)
    print(f"Wrote {OUT}")
    print(f"  estimated main-text pages: {est:.1f} "
          f"(references, exhibits and appendix excluded)")
    print("  formatting: 12-pt Times New Roman, double spaced, 1-inch margins")
    if not 5 <= est <= 15:
        print("  WARNING: estimated main text is outside the 5-15 page limit")
    print(f"  paragraphs: {len(doc.paragraphs)}   tables: {len(doc.tables)}")
